"""
Cumulative Master Excel export for the library accession register (spec
sections 6/7/11/12).

Always generated fresh from the database (never from a stored/cached
workbook -- the DB is the source of truth, spec section 8), so a repeated
download always reflects every document processed up to that moment,
including ones added since the last download, with every earlier record
still present (append-only, never overwritten).
"""
from __future__ import annotations

import calendar
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.accession_record import AccessionRecord
from app.models.document import Document

# One (header label, column width) per register column -- matches the
# reference accession-register structure (Date / Accession Number /
# Title / Organisation-Author / Language / Year / Total Pages / Type),
# not an arbitrary/invented column set.
_REGISTER_COLUMNS: list[tuple[str, int]] = [
    ("Date", 13),
    ("Accession Number", 18),
    ("Title / Book Name", 55),
    ("Organisation / Author Name", 35),
    ("Language", 16),
    ("Year of Publication", 16),
    ("Total Pages of Book", 16),
    ("Type", 14),
]
_HEADER_FILLS = [
    "FFF2CC",  # Date -- pale yellow
    "FFF2CC",  # Accession Number -- pale yellow
    "D9EAD3",  # Title / Book Name -- pale green
    "F4F4F4",  # Organisation / Author -- neutral
    "FCE5CD",  # Language -- pale orange
    "CFE2F3",  # Year of Publication -- pale blue
    "CFE2F3",  # Total Pages -- pale blue
    "F4F4F4",  # Type -- neutral
]
_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)


def render_master_register_excel(records: list[AccessionRecord], documents: list[Document] | None = None) -> bytes:
    wb = Workbook()

    common_sheet = wb.active
    common_sheet.title = "Common"
    _write_register_sheet(common_sheet, records)

    by_month: dict[tuple[int, int], list[AccessionRecord]] = defaultdict(list)
    for r in records:
        by_month[(r.record_date.year, r.record_date.month)].append(r)
    for (year, month) in sorted(by_month.keys()):
        sheet_name = f"{calendar.month_name[month]} {year}"[:31]  # Excel sheet-name length limit
        _write_register_sheet(wb.create_sheet(sheet_name), by_month[(year, month)])

    if documents is not None:
        _write_document_register_sheet(wb.create_sheet("Document Register"), documents)

    _write_summary_sheet(wb.create_sheet("Summary"), records, documents or [])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_register_sheet(sheet: Worksheet, records: list[AccessionRecord]) -> None:
    for col_idx, ((label, width), fill) in enumerate(zip(_REGISTER_COLUMNS, _HEADER_FILLS), start=1):
        cell = sheet.cell(row=1, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, record in enumerate(records, start=2):
        date_cell = sheet.cell(row=row_idx, column=1, value=record.record_date)
        date_cell.number_format = "DD-MM-YYYY"
        sheet.cell(row=row_idx, column=2, value=record.accession_number)
        sheet.cell(row=row_idx, column=3, value=record.book_name)
        sheet.cell(row=row_idx, column=4, value=record.creator)
        sheet.cell(row=row_idx, column=5, value=record.language)
        sheet.cell(row=row_idx, column=6, value=record.year_of_publication)
        sheet.cell(row=row_idx, column=7, value=record.total_pages)
        sheet.cell(row=row_idx, column=8, value=record.record_type)

    last_row = max(2, len(records) + 1)
    last_col = get_column_letter(len(_REGISTER_COLUMNS))
    sheet.freeze_panes = "A2"
    if records:
        sheet.auto_filter.ref = f"A1:{last_col}{last_row}"


def _write_document_register_sheet(sheet: Worksheet, documents: list[Document]) -> None:
    headers = [
        ("Document ID", 38), ("Filename", 45), ("Upload Date", 20), ("Pages", 10),
        ("Status", 18), ("Error", 40), ("Document Hash", 68),
    ]
    for col_idx, (label, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, doc in enumerate(documents, start=2):
        sheet.cell(row=row_idx, column=1, value=doc.id)
        sheet.cell(row=row_idx, column=2, value=doc.original_filename)
        upload_cell = sheet.cell(row=row_idx, column=3, value=doc.created_at.replace(tzinfo=None) if doc.created_at else None)
        upload_cell.number_format = "DD-MM-YYYY HH:MM"
        sheet.cell(row=row_idx, column=4, value=doc.page_count)
        sheet.cell(row=row_idx, column=5, value=doc.status.value if hasattr(doc.status, "value") else doc.status)
        sheet.cell(row=row_idx, column=6, value=doc.error_message)
        sheet.cell(row=row_idx, column=7, value=doc.document_hash)

    last_row = max(2, len(documents) + 1)
    sheet.freeze_panes = "A2"
    if documents:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"


def _write_summary_sheet(sheet: Worksheet, records: list[AccessionRecord], documents: list[Document]) -> None:
    sheet.cell(row=1, column=1, value="Master Register Summary").font = _TITLE_FONT
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 16

    completed = sum(1 for d in documents if str(getattr(d.status, "value", d.status)) == "COMPLETED")
    failed = sum(1 for d in documents if str(getattr(d.status, "value", d.status)) == "FAILED")
    needs_review = sum(1 for r in records if r.needs_review)

    by_year: dict[int, int] = defaultdict(int)
    by_month: dict[str, int] = defaultdict(int)
    for r in records:
        by_year[r.record_date.year] += 1
        by_month[f"{calendar.month_name[r.record_date.month]} {r.record_date.year}"] += 1

    row = 3
    for label, value in [
        ("Total PDFs processed", len(documents)),
        ("Successful documents", completed),
        ("Failed documents", failed),
        ("Total records extracted", len(records)),
        ("Records needing review", needs_review),
    ]:
        sheet.cell(row=row, column=1, value=label).font = _HEADER_FONT
        sheet.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Records by year").font = _HEADER_FONT
    row += 1
    for year in sorted(by_year):
        sheet.cell(row=row, column=1, value=year)
        sheet.cell(row=row, column=2, value=by_year[year])
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Records by month").font = _HEADER_FONT
    row += 1
    for month_label in sorted(by_month, key=lambda m: (int(m.split()[-1]), list(calendar.month_name).index(m.split()[0]))):
        sheet.cell(row=row, column=1, value=month_label)
        sheet.cell(row=row, column=2, value=by_month[month_label])
        row += 1
