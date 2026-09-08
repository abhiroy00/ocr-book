"""
DOCX reconstruction (spec section 18).

DOCX has no free-form absolute-positioning canvas (unlike PDF), so this
exporter reconstructs *reading-order, structural* fidelity rather than
pixel positions: real Word headings, paragraphs and tables (with actual
merged cells) in z_order, one Word section per page sized to match the
original page's aspect ratio, with a page break between pages. Complex
graphical elements (signatures/stamps/handwritten marks/photos) are
embedded as images, sized proportionally to their footprint on the
original page. This is NOT claimed to be pixel-identical to the PDF
reconstruction — see README known limitations.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Emu, Pt

from app.core.logging import get_logger
from app.models.enums import LayoutBlockType, TextAlign
from app.schemas.document_json import DocumentBlockJSON, PageJSON
from app.services.storage import get_storage
from app.utils.coordinates import page_size_emu, px_to_emu

logger = get_logger(__name__)

_IMAGE_BLOCK_TYPES = {
    LayoutBlockType.IMAGE,
    LayoutBlockType.CHART,
    LayoutBlockType.SIGNATURE,
    LayoutBlockType.STAMP,
    LayoutBlockType.HANDWRITTEN,
}

_ALIGN_MAP = {
    TextAlign.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    TextAlign.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    TextAlign.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
}


def render_document_docx(pages: list[PageJSON]) -> bytes:
    document = Document()
    storage = get_storage()

    for page_index, page_json in enumerate(pages):
        section = document.sections[0] if page_index == 0 else document.add_section()
        width_emu, height_emu = page_size_emu(page_json.page_width, page_json.page_height, page_json.dpi)
        section.page_width = Emu(width_emu)
        section.page_height = Emu(height_emu)
        section.left_margin = section.right_margin = Emu(int(0.5 * 914400))
        section.top_margin = section.bottom_margin = Emu(int(0.5 * 914400))

        for block in page_json.sorted_blocks():
            try:
                _render_block(document, block, page_json, storage)
            except Exception as exc:  # noqa: BLE001 - one bad block must not fail the export
                logger.warning("docx_block_render_failed", block_id=block.id, block_type=block.type.value, error=str(exc))

        if page_index < len(pages) - 1:
            document.add_page_break()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_block(document: Document, block: DocumentBlockJSON, page_json: PageJSON, storage) -> None:
    if block.type == LayoutBlockType.TABLE and block.table:
        _render_table(document, block)
        return

    if block.type == LayoutBlockType.HLINE:
        _render_horizontal_rule(document)
        return

    if block.type == LayoutBlockType.VLINE:
        return  # no reliable DOCX primitive for a free-standing vertical rule

    if block.type in _IMAGE_BLOCK_TYPES:
        _render_image(document, block, page_json, storage)
        return

    _render_text(document, block)


def _render_text(document: Document, block: DocumentBlockJSON) -> None:
    text = "\n".join(run.text for run in block.content) if block.content else ""
    if not text.strip():
        return

    if block.type == LayoutBlockType.TITLE:
        paragraph = document.add_heading(level=0)
    elif block.type == LayoutBlockType.HEADING:
        paragraph = document.add_heading(level=1)
    elif block.type == LayoutBlockType.SUBHEADING:
        paragraph = document.add_heading(level=2)
    else:
        paragraph = document.add_paragraph()

    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        if i < len(lines) - 1:
            run.add_break()
        font_size = block.style.get("font_size")
        if font_size:
            run.font.size = Pt(float(font_size))
        if block.type in (LayoutBlockType.FOOTNOTE, LayoutBlockType.PAGE_NUMBER, LayoutBlockType.HEADER, LayoutBlockType.FOOTER):
            run.font.size = Pt(float(font_size or 9))
            run.font.italic = block.type == LayoutBlockType.FOOTNOTE

    align_name = (block.style.get("alignment") or "LEFT").upper()
    try:
        paragraph.alignment = _ALIGN_MAP[TextAlign(align_name)]
    except (KeyError, ValueError):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _render_table(document: Document, block: DocumentBlockJSON) -> None:
    table_json = block.table
    n_rows = len(table_json.rows)
    n_cols = max((len(r.cells) for r in table_json.rows), default=0)
    if n_rows == 0 or n_cols == 0:
        return

    # Compute the actual column count from max(column + colspan) across all
    # cells, since sparse rows (post row-merge) may under-report len(cells).
    n_cols = max(
        (c.column + c.colspan for r in table_json.rows for c in r.cells), default=n_cols
    )

    table = document.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    merged: set[tuple[int, int]] = set()
    for row in table_json.rows:
        for cell in row.cells:
            if cell.row >= n_rows or cell.column >= n_cols:
                continue
            docx_cell = table.cell(cell.row, cell.column)
            docx_cell.text = cell.text

            if cell.rowspan > 1 or cell.colspan > 1:
                r_end = min(cell.row + cell.rowspan - 1, n_rows - 1)
                c_end = min(cell.column + cell.colspan - 1, n_cols - 1)
                if (r_end, c_end) != (cell.row, cell.column):
                    end_cell = table.cell(r_end, c_end)
                    docx_cell = docx_cell.merge(end_cell)

            for paragraph in docx_cell.paragraphs:
                try:
                    paragraph.alignment = _ALIGN_MAP[cell.align_h]
                except KeyError:
                    pass
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.bold = cell.is_header


def _render_image(document: Document, block: DocumentBlockJSON, page_json: PageJSON, storage) -> None:
    if not block.image_ref or not storage.exists(block.image_ref):
        return
    data = storage.read(block.image_ref)
    width_emu = px_to_emu(block.bbox.width, page_json.dpi)
    document.add_picture(io.BytesIO(data), width=Emu(width_emu))


def _render_horizontal_rule(document: Document) -> None:
    paragraph = document.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_borders.append(bottom)
    p_pr.append(p_borders)
