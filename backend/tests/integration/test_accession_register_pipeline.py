"""
Integration test for the Automatic Master Accession Register
(`app.services.accession_register_service.append_document_record`)
against a real (in-memory sqlite) database: simulates a completed
processing job the way `app.workers.pipeline_tasks._run_pipeline` leaves
one, and verifies the workbook's row count grows and the searchable PDF
filename is stored correctly.

Mirrors `tests/integration/test_accession_service.py`'s convention: real
DB + real filesystem, no live Celery/Redis broker needed.
"""
from __future__ import annotations

from datetime import date

import pytest

from openpyxl import load_workbook

from app.models.accession_record import AccessionRecord
from app.models.enums import DocumentStatus, ExportType, OCRProviderEnum, PreprocessProfileEnum
from app.services import accession_register_service as reg
from app.services import document_service


@pytest.fixture(autouse=True)
def _patch_storage(monkeypatch, tmp_storage):
    monkeypatch.setattr(reg, "get_storage", lambda: tmp_storage)
    return tmp_storage


def _simulate_completed_processing_job(
    db, storage, sample_png_bytes, accession_number, book_name, filename, pdf_filename="searchable.pdf"
):
    """Reproduces the exact end state `_run_pipeline` leaves a document in
    right before it calls `accession_register_service.append_document_record`:
    a `ProcessingJob` at `COMPLETED`, a DB-backed `AccessionRecord`, and a
    stored `SEARCHABLE_PDF` export."""
    document, job = document_service.create_document_from_upload(
        db, storage, filename, sample_png_bytes, OCRProviderEnum.PADDLEOCR, 300, PreprocessProfileEnum.BALANCED
    )
    document.page_count = 120
    db.add(AccessionRecord(
        document_id=document.id, accession_number=accession_number, book_name=book_name,
        creator="Ministry Of Testing", language="HINDI/ENGLISH", year_of_publication="1998",
        total_pages=document.page_count, record_date=date.today(),
    ))
    job.status = DocumentStatus.COMPLETED
    db.commit()
    pdf_path = f"output/{document.id}/{pdf_filename}"
    document_service.record_export_file(db, document, ExportType.SEARCHABLE_PDF, pdf_path, 2048)
    return document


def test_workbook_row_count_increases_as_each_document_completes(test_db_session, tmp_storage, sample_png_bytes):
    path = reg.ensure_master_excel()
    assert load_workbook(path)[reg.SHEET_NAME].max_row == 1  # header only, nothing processed yet

    doc_a = _simulate_completed_processing_job(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")
    assert reg.append_document_record(test_db_session, doc_a.id) is True
    assert load_workbook(path)[reg.SHEET_NAME].max_row == 2

    doc_b = _simulate_completed_processing_job(test_db_session, tmp_storage, sample_png_bytes, "D-2", "Book B", "b.png")
    assert reg.append_document_record(test_db_session, doc_b.id) is True
    assert load_workbook(path)[reg.SHEET_NAME].max_row == 3

    doc_c = _simulate_completed_processing_job(test_db_session, tmp_storage, sample_png_bytes, "D-3", "Book C", "c.png")
    assert reg.append_document_record(test_db_session, doc_c.id) is True
    assert load_workbook(path)[reg.SHEET_NAME].max_row == 4


def test_searchable_pdf_filename_and_path_stored_correctly(test_db_session, tmp_storage, sample_png_bytes):
    document = _simulate_completed_processing_job(
        test_db_session, tmp_storage, sample_png_bytes, "D-1", "Statistical Abstract", "book.png",
        pdf_filename="searchable.pdf",
    )
    assert reg.append_document_record(test_db_session, document.id) is True

    sheet = load_workbook(reg.ensure_master_excel())[reg.SHEET_NAME]
    assert sheet.cell(row=2, column=reg.COL_PDF_FILENAME).value == "searchable.pdf"
    assert sheet.cell(row=2, column=reg.COL_PDF_PATH).value == f"output/{document.id}/searchable.pdf"


def test_full_row_reflects_the_accession_record_and_export(test_db_session, tmp_storage, sample_png_bytes):
    document = _simulate_completed_processing_job(
        test_db_session, tmp_storage, sample_png_bytes, "D-42", "STATISTICAL ABSTRACT OF PUNJAB 1998", "book.png",
    )
    assert reg.append_document_record(test_db_session, document.id) is True

    sheet = load_workbook(reg.ensure_master_excel())[reg.SHEET_NAME]
    # openpyxl always reads a date-formatted cell back as `datetime` (even
    # though a plain `date` was written) -- see the same note in
    # `tests/unit/test_master_register_exporter.py`.
    assert sheet.cell(row=2, column=reg.COL_DATE).value.date() == date.today()
    assert sheet.cell(row=2, column=reg.COL_ACCESSION).value == "D-42"
    assert sheet.cell(row=2, column=reg.COL_BOOK_NAME).value == "STATISTICAL ABSTRACT OF PUNJAB 1998"
    assert sheet.cell(row=2, column=reg.COL_PAGE).value == 120
    assert sheet.cell(row=2, column=reg.COL_LANGUAGE).value == "HINDI/ENGLISH"
    assert sheet.cell(row=2, column=reg.COL_YEAR).value == "1998"
    assert sheet.cell(row=2, column=reg.COL_CREATOR).value == "Ministry Of Testing"
    assert sheet.cell(row=2, column=reg.COL_TIMESTAMP).value is not None
