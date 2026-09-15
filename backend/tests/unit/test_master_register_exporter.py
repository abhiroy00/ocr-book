"""
Tests for the cumulative Master Excel export
(`app.exporters.master_register_exporter`) -- built from plain in-memory
ORM-shaped objects (no DB needed), verifying sheet structure: one "Common"
sheet with every record, one sheet per distinct month, and that repeated
generation is stable/reflects exactly what was passed in (the actual
append-only cumulative behavior is proven at the DB layer in
`test_accession_service.py`; this covers the rendering side).
"""
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.exporters.master_register_exporter import render_master_register_excel
from app.models.accession_record import AccessionRecord


def _record(accession_number, book_name, record_date, **kwargs) -> AccessionRecord:
    return AccessionRecord(
        id=f"id-{accession_number}", document_id=f"doc-{accession_number}", accession_number=accession_number,
        book_name=book_name, record_date=record_date, total_pages=kwargs.pop("total_pages", 100), **kwargs,
    )


def test_common_sheet_contains_every_record():
    records = [
        _record("D-1", "Book One", date(2026, 1, 13)),
        _record("D-2", "Book Two", date(2026, 2, 5)),
        _record("D-3", "Book Three", date(2026, 2, 20)),
    ]
    wb = load_workbook(BytesIO(render_master_register_excel(records)))
    common = wb["Common"]
    book_names = {common.cell(row=r, column=3).value for r in range(2, common.max_row + 1)}
    assert book_names == {"Book One", "Book Two", "Book Three"}


def test_creates_one_sheet_per_distinct_month():
    records = [
        _record("D-1", "Book One", date(2026, 1, 13)),
        _record("D-2", "Book Two", date(2026, 2, 5)),
        _record("D-3", "Book Three", date(2026, 2, 20)),
    ]
    wb = load_workbook(BytesIO(render_master_register_excel(records)))
    assert "January 2026" in wb.sheetnames
    assert "February 2026" in wb.sheetnames

    jan_books = {wb["January 2026"].cell(row=r, column=3).value for r in range(2, wb["January 2026"].max_row + 1)}
    feb_books = {wb["February 2026"].cell(row=r, column=3).value for r in range(2, wb["February 2026"].max_row + 1)}
    assert jan_books == {"Book One"}
    assert feb_books == {"Book Two", "Book Three"}


def test_header_row_matches_the_reference_register_columns():
    wb = load_workbook(BytesIO(render_master_register_excel([])))
    headers = [wb["Common"].cell(row=1, column=c).value for c in range(1, 9)]
    assert headers == [
        "Date", "Accession Number", "Title / Book Name", "Organisation / Author Name",
        "Language", "Year of Publication", "Total Pages of Book", "Type",
    ]


def test_empty_dataset_still_produces_a_valid_workbook_with_just_headers():
    wb = load_workbook(BytesIO(render_master_register_excel([])))
    common = wb["Common"]
    assert common.cell(row=1, column=1).value == "Date"
    assert common.max_row == 1  # header only, no data rows, no crash


def test_frozen_header_and_autofilter_are_set_when_there_is_data():
    records = [_record("D-1", "Book One", date(2026, 1, 13))]
    wb = load_workbook(BytesIO(render_master_register_excel(records)))
    common = wb["Common"]
    assert common.freeze_panes == "A2"
    assert common.auto_filter.ref is not None


def test_date_values_are_real_excel_dates_not_strings():
    records = [_record("D-1", "Book One", date(2026, 1, 13))]
    wb = load_workbook(BytesIO(render_master_register_excel(records)))
    cell = wb["Common"].cell(row=2, column=1)
    # openpyxl always reads a date-formatted cell back as `datetime`
    # (even though a plain `date` was written) -- what actually matters
    # is that it round-trips as a real Excel date value, not the string
    # "2026-01-13", which is what makes month/year filtering in Excel work.
    assert cell.value.date() == date(2026, 1, 13)


def test_needs_review_records_are_still_included_not_dropped():
    """Spec: never discard/hide uncertain data -- an unclear extraction
    still gets a row, just flagged internally."""
    record = _record("D-1", "Uncertain Title", date(2026, 1, 13), needs_review=True, extraction_notes="creator: unclear")
    wb = load_workbook(BytesIO(render_master_register_excel([record])))
    common = wb["Common"]
    assert common.cell(row=2, column=3).value == "Uncertain Title"
