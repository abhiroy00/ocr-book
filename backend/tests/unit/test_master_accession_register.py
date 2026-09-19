"""
Tests for the Automatic Master Accession Register
(`app.services.accession_register_service`) -- workbook creation/
formatting plus the `append_document_record` entry point's own
preconditions: only a genuinely COMPLETED job with a searchable PDF gets
a row, duplicates (by accession number) are skipped, and concurrent
writers never corrupt the file.

`get_storage()` is monkeypatched to a temp directory (via the shared
`tmp_storage` fixture) so nothing here ever touches the real project
`storage/` folder.
"""
from __future__ import annotations

import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from app.models.accession_record import AccessionRecord
from app.models.enums import DocumentStatus, ExportType, OCRProviderEnum, PreprocessProfileEnum
from app.services import accession_register_service as reg
from app.services import document_service


@pytest.fixture(autouse=True)
def _patch_storage(monkeypatch, tmp_storage):
    monkeypatch.setattr(reg, "get_storage", lambda: tmp_storage)
    return tmp_storage


def _make_completed_document(db, storage, sample_png_bytes, accession_number, book_name="Book", filename="book.png"):
    """A document whose (single) `ProcessingJob` is COMPLETED, with an
    `AccessionRecord` and a `SEARCHABLE_PDF` export already in place --
    exactly the state the real pipeline leaves a document in right before
    calling `append_document_record`."""
    document, job = document_service.create_document_from_upload(
        db, storage, filename, sample_png_bytes, OCRProviderEnum.PADDLEOCR, 300, PreprocessProfileEnum.BALANCED
    )
    document.page_count = 42
    job.status = DocumentStatus.COMPLETED
    db.add(AccessionRecord(
        document_id=document.id, accession_number=accession_number, book_name=book_name,
        creator="Ministry Of Testing", language="ENGLISH", year_of_publication="2020",
        total_pages=document.page_count, record_date=date.today(),
    ))
    db.commit()
    document_service.record_export_file(db, document, ExportType.SEARCHABLE_PDF, f"output/{document.id}/searchable.pdf", 1024)
    return document, job


def _register_sheet(path):
    return load_workbook(path)[reg.SHEET_NAME]


# ---------------------------------------------------------------------------
# Workbook created / header created
# ---------------------------------------------------------------------------

def test_workbook_created(tmp_storage):
    path = reg.ensure_master_excel()
    assert tmp_storage.exists(reg.REGISTER_RELATIVE_PATH)
    assert load_workbook(path)[reg.SHEET_NAME] is not None


def test_header_created(tmp_storage):
    sheet = _register_sheet(reg.ensure_master_excel())
    assert [sheet.cell(row=1, column=c).value for c in range(1, len(reg.HEADERS) + 1)] == reg.HEADERS

    header_cell = sheet.cell(row=1, column=1)
    assert header_cell.fill.start_color.rgb in ("003A8E2D", "3A8E2D")
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb in ("00FFFFFF", "FFFFFF")
    assert header_cell.alignment.horizontal == "center"
    assert header_cell.border.top.style == "thin"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


# ---------------------------------------------------------------------------
# Append / duplicate / failed-job behavior
# ---------------------------------------------------------------------------

def test_first_row_appended(test_db_session, tmp_storage, sample_png_bytes):
    document, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", book_name="Book One")

    assert reg.append_document_record(test_db_session, document.id) is True

    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=reg.COL_ACCESSION).value == "D-1"
    assert sheet.cell(row=2, column=reg.COL_BOOK_NAME).value == "Book One"
    assert sheet.cell(row=2, column=reg.COL_PDF_FILENAME).value == "searchable.pdf"
    assert sheet.cell(row=2, column=reg.COL_PDF_PATH).value == f"output/{document.id}/searchable.pdf"
    # Data-row formatting: wrapped, top-aligned (spec), distinct from the header's own alignment.
    assert sheet.cell(row=2, column=reg.COL_BOOK_NAME).alignment.wrap_text is True
    assert sheet.cell(row=2, column=reg.COL_BOOK_NAME).alignment.vertical == "top"


def test_second_row_appended(test_db_session, tmp_storage, sample_png_bytes):
    doc_a, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")
    doc_b, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-2", "Book B", "b.png")

    assert reg.append_document_record(test_db_session, doc_a.id) is True
    assert reg.append_document_record(test_db_session, doc_b.id) is True

    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == 3
    book_names = {sheet.cell(row=r, column=reg.COL_BOOK_NAME).value for r in (2, 3)}
    assert book_names == {"Book A", "Book B"}


def test_duplicate_accession_number_is_skipped(test_db_session, tmp_storage, sample_png_bytes):
    document, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-305", "Book One")

    assert reg.append_document_record(test_db_session, document.id) is True
    assert reg.append_document_record(test_db_session, document.id) is False  # same accession number, second call

    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == 2  # still just the one row


def test_failed_processing_does_not_append(test_db_session, tmp_storage, sample_png_bytes):
    document, job = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book One")
    job.status = DocumentStatus.FAILED
    test_db_session.commit()

    assert reg.append_document_record(test_db_session, document.id) is False

    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == 1  # header only -- nothing appended


def test_cancelled_processing_does_not_append(test_db_session, tmp_storage, sample_png_bytes):
    document, job = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book One")
    job.status = DocumentStatus.CANCELLED
    test_db_session.commit()

    assert reg.append_document_record(test_db_session, document.id) is False


def test_missing_searchable_pdf_does_not_append(test_db_session, tmp_storage, sample_png_bytes):
    document, job = document_service.create_document_from_upload(
        test_db_session, tmp_storage, "a.png", sample_png_bytes, OCRProviderEnum.PADDLEOCR, 300, PreprocessProfileEnum.BALANCED
    )
    job.status = DocumentStatus.COMPLETED
    test_db_session.add(AccessionRecord(
        document_id=document.id, accession_number="D-1", book_name="Book One",
        total_pages=document.page_count, record_date=date.today(),
    ))
    test_db_session.commit()
    # No SEARCHABLE_PDF export recorded -- the job says COMPLETED, but the
    # spec's second precondition ("AND searchable_pdf exists") isn't met.

    assert reg.append_document_record(test_db_session, document.id) is False


# ---------------------------------------------------------------------------
# Concurrency
# ---------------------------------------------------------------------------

def test_concurrent_append_does_not_corrupt_workbook(test_db_session, tmp_storage, sample_png_bytes):
    engine = test_db_session.get_bind()
    SessionFactory = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    documents = [
        _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, f"D-{i + 1}", f"Book {i + 1}", f"book_{i}.png")[0]
        for i in range(6)
    ]

    def _append(document_id: str) -> bool:
        session = SessionFactory()
        try:
            return reg.append_document_record(session, document_id)
        finally:
            session.close()

    with ThreadPoolExecutor(max_workers=6) as executor:
        results = list(executor.map(_append, [d.id for d in documents]))

    assert all(results)
    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == len(documents) + 1
    accessions = {sheet.cell(row=r, column=reg.COL_ACCESSION).value for r in range(2, sheet.max_row + 1)}
    assert accessions == {f"D-{i + 1}" for i in range(6)}


def test_concurrent_lock_holders_never_overlap(tmp_storage):
    reg.ensure_master_excel()
    overlap_detected = threading.Event()
    active = 0
    guard = threading.Lock()

    def worker():
        nonlocal active
        lock = reg.acquire_lock()
        try:
            with guard:
                active += 1
                if active > 1:
                    overlap_detected.set()
        finally:
            with guard:
                active -= 1
            reg.release_lock(lock)

    threads = [threading.Thread(target=worker) for _ in range(5)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=10)

    assert not overlap_detected.is_set()


# ---------------------------------------------------------------------------
# Backfill (books completed before the register existed / failed live appends)
# and the download route
# ---------------------------------------------------------------------------

def test_sync_backfills_completed_documents_missing_from_register(test_db_session, tmp_storage, sample_png_bytes):
    for i in range(3):
        _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, f"D-{i + 1}", f"Book {i + 1}", f"b{i}.png")
    failed_doc, failed_job = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-9", "Failed", "f.png")
    failed_job.status = DocumentStatus.FAILED
    test_db_session.commit()

    assert reg.sync_register_from_db(test_db_session) == 3

    sheet = _register_sheet(reg.ensure_master_excel())
    assert sheet.max_row == 4  # header + the 3 completed books; the FAILED one is not backfilled
    accessions = {sheet.cell(row=r, column=reg.COL_ACCESSION).value for r in range(2, 5)}
    assert accessions == {"D-1", "D-2", "D-3"}


def test_sync_is_idempotent_and_skips_rows_already_present(test_db_session, tmp_storage, sample_png_bytes):
    doc_a, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")
    _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-2", "Book B", "b.png")
    assert reg.append_document_record(test_db_session, doc_a.id) is True  # A already in the register

    assert reg.sync_register_from_db(test_db_session) == 1  # only B is missing
    assert reg.sync_register_from_db(test_db_session) == 0  # nothing left to add

    assert _register_sheet(reg.ensure_master_excel()).max_row == 3


def test_sync_ignores_soft_deleted_documents(test_db_session, tmp_storage, sample_png_bytes):
    document, _ = _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")
    document.is_deleted = True
    test_db_session.commit()

    assert reg.sync_register_from_db(test_db_session) == 0


def test_get_master_register_bytes_returns_the_full_up_to_date_workbook(test_db_session, tmp_storage, sample_png_bytes):
    from io import BytesIO

    _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")
    _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-2", "Book B", "b.png")

    sheet = load_workbook(BytesIO(reg.get_master_register_bytes(test_db_session)))[reg.SHEET_NAME]
    assert sheet.max_row == 3
    assert [sheet.cell(row=1, column=c).value for c in range(1, len(reg.HEADERS) + 1)] == reg.HEADERS


def test_master_register_download_route(test_db_session, tmp_storage, sample_png_bytes):
    from io import BytesIO

    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.deps import get_db
    from app.api.v1 import accession as accession_api

    _make_completed_document(test_db_session, tmp_storage, sample_png_bytes, "D-1", "Book A", "a.png")

    app = FastAPI()
    app.include_router(accession_api.router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: test_db_session

    response = TestClient(app).get("/api/accession-records/export/master-register")

    assert response.status_code == 200
    assert "master_accession_register.xlsx" in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content))[reg.SHEET_NAME]
    assert sheet.cell(row=2, column=reg.COL_ACCESSION).value == "D-1"
